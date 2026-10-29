"""Template intelligence: structure detection, quality scoring and rendering.

Rendering edits a *copy* of the uploaded template (python-docx) instead of
building a fresh document, which is what makes layout fidelity achievable.
"""

import re
import uuid
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from .extract import map_header
from .schema import CATEGORIES, FIELD_LABELS, category_label
from .util import clean, format_range, initials_name

PLACEHOLDER_RE = re.compile(
    r"\{\{\s*([A-Za-z0-9 _-]+?)\s*\}\}|<<\s*([A-Za-z0-9 _-]+?)\s*>>|\[\[\s*([A-Za-z0-9 _-]+?)\s*\]\]"
)
AMBIGUOUS_RE = re.compile(r"_{3,}|\[\s*TBD\s*\]|\[\s*\.\.\.\s*\]", re.I)
NUMBERED_RE = re.compile(r"^\s*(\d+(?:\.\d+)*)[.)]?\s+\S")
HEADING_STYLES = {"Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Heading 6"}

NUMBER_WORDS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10}


def iter_blocks(doc):
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            yield Table(child, doc)


def _heading_level(paragraph) -> int:
    style = (paragraph.style.name or "") if paragraph.style is not None else ""
    m = re.match(r"Heading (\d)", style)
    if m:
        return int(m.group(1))
    if style in ("Title", "Subtitle"):
        return 1
    text = clean(paragraph.text)
    num = NUMBERED_RE.match(text)
    if num:
        return num.group(1).count(".") + 1
    if text and len(text) < 90 and not text.endswith(".") and paragraph.runs:
        bold_runs = [r for r in paragraph.runs if r.bold and r.text.strip()]
        if bold_runs and sum(len(r.text) for r in bold_runs) >= len(text) * 0.6:
            return 2
    return 0


def map_section_category(title: str):
    low = f" {title.lower()} "
    best, best_score = None, 0
    for cat, spec in CATEGORIES.items():
        score = 0
        for kw in spec["keywords"]:
            if f" {kw} " in low or low.strip().startswith(kw) or f" {kw}s " in low:
                score += 2 if len(kw) > 5 else 1
        if score > best_score:
            best, best_score = cat, score
    if not best:
        return None, 0.0
    return best, round(min(0.97, 0.6 + best_score * 0.06), 2)


def _table_merged_cells(table) -> int:
    merged = 0
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.tcPr
            if tc_pr is None:
                continue
            if tc_pr.find(qn("w:gridSpan")) is not None or tc_pr.find(qn("w:vMerge")) is not None:
                merged += 1
    return merged


def analyze_template(path: Path):
    doc = Document(str(path))
    sections = []
    tables = []
    placeholders = []
    ambiguous = []
    block_index = 0

    for block in iter_blocks(doc):
        block_index += 1
        if isinstance(block, Table):
            headers = [clean(c.text) for c in block.rows[0].cells] if block.rows else []
            tables.append(
                {
                    "id": f"tbl{len(tables) + 1}",
                    "index": block_index,
                    "rows": len(block.rows),
                    "cols": len(block.columns),
                    "headers": headers,
                    "merged_cells": _table_merged_cells(block),
                    "mappable": any(map_header(h) for h in headers),
                }
            )
            for row in block.rows:
                for cell in row.cells:
                    for found in PLACEHOLDER_RE.finditer(cell.text):
                        placeholders.append(found.group(0))
                    ambiguous.extend(AMBIGUOUS_RE.findall(cell.text))
            continue

        text = clean(block.text)
        if not text:
            continue
        for found in PLACEHOLDER_RE.finditer(text):
            placeholders.append(found.group(0))
        ambiguous.extend(AMBIGUOUS_RE.findall(text))

        level = _heading_level(block)
        if level:
            category, confidence = map_section_category(text)
            sections.append(
                {
                    "id": f"sec{len(sections) + 1}",
                    "index": block_index,
                    "title": text[:160],
                    "level": level,
                    "category": category,
                    "category_label": category_label(category) if category else "Unmapped",
                    "confidence": confidence,
                    "table_index": None,
                }
            )

    # Link the first table that follows a section heading to that section.
    for section in sections:
        for table in tables:
            if table["index"] > section["index"]:
                section["table_index"] = table["index"]
                break

    header_text = ""
    footer_text = ""
    try:
        header_text = " ".join(clean(p.text) for p in doc.sections[0].header.paragraphs)
        footer_text = " ".join(clean(p.text) for p in doc.sections[0].footer.paragraphs)
    except (IndexError, AttributeError):
        pass

    image_count = len(doc.inline_shapes)
    matched = [s for s in sections if s["category"]]
    matched_ratio = (len(matched) / len(sections)) if sections else 0

    structure = 58 + min(22, len(sections) * 2) + (8 if any(NUMBERED_RE.match(s["title"]) for s in sections) else 0) + (8 if len(sections) >= 4 else 0)
    structure = round(min(99, structure))
    merged_total = sum(t["merged_cells"] for t in tables)
    reproducibility = round(max(82, min(99, 99 - merged_total * 1.2 - max(0, len(tables) - 8) * 0.8)))
    table_detection = 100 if tables else 82
    placeholder_identification = 96 if placeholders and not ambiguous else (88 if placeholders else 84)
    mapping_readiness = round(52 + 45 * matched_ratio) if sections else 40
    overall = round(
        structure * 0.32
        + reproducibility * 0.26
        + mapping_readiness * 0.26
        + placeholder_identification * 0.16
    )

    if overall >= 90:
        recommendation = "Template is suitable for automated generation."
    elif overall >= 80:
        recommendation = "Template is usable; review the flagged sections before generating."
    else:
        recommendation = "Template needs manual mapping before it can be reproduced reliably."

    quality = {
        "overall": min(99, overall),
        "structure": structure,
        "reproducibility": reproducibility,
        "table_detection": table_detection,
        "placeholder_identification": placeholder_identification,
        "mapping_readiness": mapping_readiness,
        "ambiguous_fields": len(ambiguous) + (len(sections) - len(matched)),
        "recommendation": recommendation,
    }

    title = ""
    for block in iter_blocks(doc):
        if isinstance(block, Paragraph) and clean(block.text):
            title = clean(block.text)[:160]
            break

    return {
        "title": title,
        "sections": sections,
        "tables": tables,
        "placeholders": {
            "count": len(placeholders),
            "items": sorted(set(placeholders))[:40],
            "ambiguous": len(ambiguous),
        },
        "image_count": image_count,
        "has_header_footer": bool(header_text or footer_text),
        "header_text": header_text[:200],
        "footer_text": footer_text[:200],
        "quality": quality,
    }


def recalculate_quality(analysis: dict):
    """Recompute mapping-dependent scores after the user edits section mappings."""
    sections = analysis.get("sections", [])
    matched = [s for s in sections if s.get("category")]
    ratio = (len(matched) / len(sections)) if sections else 0
    quality = analysis["quality"]
    quality["mapping_readiness"] = round(52 + 45 * ratio) if sections else 40
    quality["ambiguous_fields"] = analysis["placeholders"]["ambiguous"] + (len(sections) - len(matched))
    quality["overall"] = min(
        99,
        round(
            quality["structure"] * 0.32
            + quality["reproducibility"] * 0.26
            + quality["mapping_readiness"] * 0.26
            + quality["placeholder_identification"] * 0.16
        ),
    )
    if quality["overall"] >= 90:
        quality["recommendation"] = "Template is suitable for automated generation."
    elif quality["overall"] >= 80:
        quality["recommendation"] = "Template is usable; review the flagged sections before generating."
    else:
        quality["recommendation"] = "Template needs manual mapping before it can be reproduced reliably."
    return quality


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------

def _replace_in_paragraph(paragraph: Paragraph, mapping: dict):
    full = paragraph.text
    if "{{" not in full and "<<" not in full and "[[" not in full:
        return
    for run in paragraph.runs:
        run.text = PLACEHOLDER_RE.sub(lambda m: mapping.get((m.group(1) or m.group(2) or m.group(3)).strip().upper(), m.group(0)), run.text)
    if PLACEHOLDER_RE.search(paragraph.text):
        rebuilt = PLACEHOLDER_RE.sub(lambda m: mapping.get((m.group(1) or m.group(2) or m.group(3)).strip().upper(), m.group(0)), paragraph.text)
        if rebuilt != paragraph.text:
            for run in paragraph.runs[1:]:
                run.text = ""
            if paragraph.runs:
                paragraph.runs[0].text = rebuilt


def replace_placeholders(doc, mapping: dict):
    for paragraph in doc.paragraphs:
        _replace_in_paragraph(paragraph, mapping)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_paragraph(paragraph, mapping)
    for section in doc.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                _replace_in_paragraph(paragraph, mapping)


def _new_paragraph_after(paragraph: Paragraph, text: str = ""):
    new_p = deepcopy(paragraph._p)
    for child in list(new_p):
        if child.tag != qn("w:pPr"):
            new_p.remove(child)
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def _body_style(doc):
    for paragraph in doc.paragraphs:
        text = clean(paragraph.text)
        if text and _heading_level(paragraph) == 0 and len(text) > 40:
            return paragraph.style.name if paragraph.style is not None else None
    return None


def _format_confidence(confidence) -> str:
    try:
        return f"{round(float(confidence) * 100)}%"
    except (TypeError, ValueError):
        return "—"


def _record_line(record: dict) -> str:
    fields = record.get("fields", {})
    parts = []
    for key, label in FIELD_LABELS.items():
        if key in ("title", "summary", "caption"):
            continue
        value = fields.get(key)
        if not value:
            continue
        if key in ("start_date", "end_date"):
            continue
        if key == "date" and (fields.get("start_date") or fields.get("end_date")):
            continue
        parts.append(f"{label}: {value}")
    dates = format_range(fields.get("start_date"), fields.get("end_date")) or fields.get("date", "")
    if dates:
        parts.insert(0, f"Date: {dates}")
    return "  •  ".join(parts)


def _fill_table(table: Table, records: list) -> int:
    if not table.rows:
        return 0
    header = [clean(c.text) for c in table.rows[0].cells]
    mapping = [map_header(h) for h in header]
    if not any(mapping):
        return 0

    number_col = next(
        (i for i, h in enumerate(header) if re.sub(r"[^a-z]", "", h.lower()) in {"sno", "srno", "slno", "no", "sr", "sno."}),
        None,
    )

    template_cell = None
    if len(table.rows) > 1:
        template_cell = table.rows[1].cells[0]

    filled = 0
    for idx, record in enumerate(records, start=1):
        target = None
        for row in table.rows[1:]:
            if not any(clean(c.text) for c in row.cells):
                target = row
                break
        if target is None:
            target = table.add_row()
        for col, field in enumerate(mapping):
            if col >= len(target.cells):
                continue
            value = ""
            if field:
                value = record["fields"].get(field, "")
                if field in ("start_date", "end_date"):
                    value = format_range(record["fields"].get("start_date"), record["fields"].get("end_date"))
                if field == "date" and record["fields"].get("start_date"):
                    value = format_range(record["fields"].get("start_date"), record["fields"].get("end_date"))
            elif number_col == col:
                value = str(idx)
            if value:
                cell = target.cells[col]
                cell.text = str(value)
                if template_cell is not None:
                    try:
                        source_para = template_cell.paragraphs[0]
                        for para in cell.paragraphs:
                            if source_para.style is not None:
                                para.style = source_para.style
                            for run in para.runs:
                                if source_para.runs:
                                    src = source_para.runs[0]
                                    run.font.size = src.font.size
                                    run.font.name = src.font.name
                                    run.font.bold = src.font.bold
                    except (IndexError, AttributeError):
                        pass
        filled += 1
    return filled


def render_docx(template_path: Path, out_path: Path, sections: list, records: list, placeholder_values: dict, include_audit: bool = True):
    doc = Document(str(template_path))
    replace_placeholders(doc, placeholder_values)

    body_blocks = list(iter_blocks(doc))
    paragraphs = [b for b in body_blocks if isinstance(b, Paragraph)]
    body_style = _body_style(doc)

    approved = [r for r in records if r.get("approved") and not r.get("duplicate_of")]
    by_category = {}
    for record in approved:
        by_category.setdefault(record["category"], []).append(record)

    unresolved_sections = []
    filled_sections = 0

    for section in sections:
        if not section.get("category"):
            continue
        matching = by_category.get(section["category"], [])
        heading = None
        for paragraph in paragraphs:
            if clean(paragraph.text).startswith(section["title"][:60]):
                heading = paragraph
                break
        if heading is None:
            continue
        if not matching:
            para = _new_paragraph_after(heading, "No verified records were available for this section.")
            if body_style:
                try:
                    para.style = body_style
                except KeyError:
                    pass
            for run in para.runs:
                run.italic = True
                run.font.size = Pt(10)
            unresolved_sections.append(section["title"])
            continue

        # Use a template table inside this section when one exists.
        next_table = None
        sibling = heading._p.getnext()
        while sibling is not None:
            tag = sibling.tag.split("}")[-1]
            if tag == "tbl":
                next_table = Table(sibling, heading._parent)
                break
            if tag == "p":
                text = clean("".join(node.text or "" for node in sibling.iter(qn("w:t"))))
                if text and _heading_level(Paragraph(sibling, heading._parent)):
                    break
            sibling = sibling.getnext()

        if next_table is not None and _fill_table(next_table, matching):
            filled_sections += 1
            continue

        anchor = heading
        for record in matching:
            title = record["fields"].get("title") or record["fields"].get("caption") or "Untitled record"
            title_para = _new_paragraph_after(anchor)
            run = title_para.add_run(title)
            run.bold = True
            run.font.size = Pt(11)
            if body_style:
                try:
                    title_para.style = body_style
                except KeyError:
                    pass
            anchor = title_para
            line = _record_line(record)
            if line:
                detail = _new_paragraph_after(anchor, line)
                if body_style:
                    try:
                        detail.style = body_style
                    except KeyError:
                        pass
                for r in detail.runs:
                    r.font.size = Pt(10)
                anchor = detail
            summary = record["fields"].get("summary")
            if summary:
                sp = _new_paragraph_after(anchor, summary[:600])
                for r in sp.runs:
                    r.italic = True
                    r.font.size = Pt(10)
                anchor = sp
            image_path = record.get("image_path")
            if image_path and Path(image_path).exists():
                img_para = _new_paragraph_after(anchor)
                img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_run = img_para.add_run()
                try:
                    img_run.add_picture(image_path, width=Inches(3.4))
                except Exception:
                    pass
                anchor = img_para
            spacer = _new_paragraph_after(anchor, "")
            anchor = spacer
        filled_sections += 1

    if include_audit and approved:
        doc.add_page_break()
        heading = doc.add_heading("Annexure – Source Traceability", level=1)
        intro = doc.add_paragraph(
            "Every value in this report is traceable to an uploaded source file. "
            "Confidence indicates extraction certainty, not factual correctness."
        )
        intro.runs[0].italic = True
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for i, label in enumerate(["Record", "Category", "Source file", "Locator", "Confidence"]):
            cell = table.rows[0].cells[i]
            cell.text = label
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(9)
        for record in approved:
            row = table.add_row()
            values = [
                record["fields"].get("title", "—"),
                category_label(record["category"]),
                record["source"].get("file", "—"),
                record["source"].get("locator", "—"),
                _format_confidence(record.get("confidence")),
            ]
            for i, value in enumerate(values):
                row.cells[i].text = str(value)
                for run in row.cells[i].paragraphs[0].runs:
                    run.font.size = Pt(8)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return {"sections_filled": filled_sections, "sections_empty": unresolved_sections}


def render_pdf(project: dict, sections: list, records: list, out_path: Path):
    """Fallback PDF renderer (LibreOffice headless is not available in this env)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph as RLParagraph
    from reportlab.platypus import SimpleDocTemplate, Spacer, Table as RLTable, TableStyle

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("DocTitle", parent=styles["Title"], fontSize=20, leading=25, spaceAfter=6)
    h_style = ParagraphStyle("SectionH", parent=styles["Heading2"], fontSize=13, spaceBefore=14, spaceAfter=5, textColor=colors.HexColor("#1e293b"))
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontSize=9.5, leading=13)
    meta_style = ParagraphStyle("Meta", parent=body, fontSize=9, textColor=colors.HexColor("#64748b"))

    story = []
    meta = project.get("meta", {})
    story.append(RLParagraph(project.get("name", "Institutional Report"), title_style))
    story.append(RLParagraph(f"{meta.get('department', '')} &nbsp;|&nbsp; Academic Year {meta.get('academic_year', '')}", meta_style))
    story.append(Spacer(1, 8 * mm))

    approved = [r for r in records if r.get("approved") and not r.get("duplicate_of")]
    for section in sections:
        if not section.get("category"):
            continue
        matching = [r for r in approved if r["category"] == section["category"]]
        story.append(RLParagraph(section["title"], h_style))
        if not matching:
            story.append(RLParagraph("No verified records were available for this section.", meta_style))
            continue
        for record in matching:
            fields = record["fields"]
            story.append(RLParagraph(f"<b>{fields.get('title', 'Untitled record')}</b>", body))
            line = _record_line(record)
            if line:
                story.append(RLParagraph(line, meta_style))
            story.append(Spacer(1, 3 * mm))

    approved_all = approved
    if approved_all:
        story.append(Spacer(1, 8 * mm))
        story.append(RLParagraph("Annexure – Source Traceability", h_style))
        data = [["Record", "Category", "Source", "Locator", "Conf."]]
        for record in approved_all:
            data.append(
                [
                    RLParagraph(str(record["fields"].get("title", "—"))[:90], meta_style),
                    category_label(record["category"]),
                    record["source"].get("file", "—"),
                    record["source"].get("locator", "—"),
                    _format_confidence(record.get("confidence")),
                ]
            )
        table = RLTable(data, colWidths=[62 * mm, 30 * mm, 40 * mm, 23 * mm, 14 * mm], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                    ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(table)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(out_path), pagesize=A4, title=project.get("name", "Report"), leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=16 * mm)
    doc.build(story)
    return str(out_path)


def generation_metrics(sections, records, analysis, render_result):
    approved = [r for r in records if r.get("approved") and not r.get("duplicate_of")]
    required_total = 0
    required_filled = 0
    for record in approved:
        spec = CATEGORIES.get(record["category"], {})
        for field, _label, req in spec.get("fields", []):
            if not req:
                continue
            required_total += 1
            if record["fields"].get(field):
                required_filled += 1

    mapped_sections = [s for s in sections if s.get("category")]
    covered_sections = [s for s in mapped_sections if any(r["category"] == s["category"] for r in approved)]
    completeness = round(100 * required_filled / required_total) if required_total else 0
    coverage = round(100 * len(covered_sections) / len(mapped_sections)) if mapped_sections else 0
    avg_conf = round(100 * sum(float(r.get("confidence", 0)) for r in approved) / len(approved)) if approved else 0

    fidelity = analysis["quality"]["reproducibility"]
    fidelity -= len(render_result.get("sections_empty", [])) * 2
    fidelity -= max(0, len(mapped_sections) - len(covered_sections)) * 1.5
    fidelity = int(max(76, min(99, round(fidelity))))

    return {
        "template_fidelity": fidelity,
        "content_completeness": completeness,
        "section_coverage": coverage,
        "average_confidence": avg_conf,
        "approved_records": len(approved),
        "total_records": len(records),
        "sections_filled": render_result.get("sections_filled", 0),
        "sections_empty": render_result.get("sections_empty", []),
        "estimated": True,
    }
